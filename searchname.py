# -*- coding: utf-8 -*-

# Form implementation generated from reading ui file 'searchname.ui',
# licensing of 'searchname.ui' applies.
#
# Created: Sat Mar 14 20:44:19 2020
#      by: pyside2-uic  running on PySide2 5.13.0
#
# WARNING! All changes made in this file will be lost!

from PySide2 import QtCore, QtGui, QtWidgets
from PyQt5.QtWidgets import *
from PyQt5.QtPrintSupport import QPrinter, QPrintDialog, QPrintPreviewDialog
from PyQt5.QtGui import *
import search
import info
import os
import docx
from docx import Document
import webbrowser
from docx.shared import Inches
from fpdf import FPDF
import PyPDF2
from io import BytesIO
import PIL.Image
from PIL import  Image
import codecs
import time
import sqlite3
conn = sqlite3.connect('database.db')

class Ui_Dialog(object):
    def openWindowtxt(self, Dialogcolor):
        searchrol = str(self.lineEdit.text())
        cur = conn.cursor()
        cur.execute("SELECT * from artefacts WHERE id_artefact=?", (searchrol,))
        dat = cur.fetchall()
        if dat:
            f = open('artefact.txt', 'w')
            for data in dat:
                f.write("Отчет по артефакту, где ID " + data[2]+"\n")
                f.write("\nID артефакта: " + data[2])
                f.write("\nТип артефакта: " + data[5])
                f.write("\nЦвет артефакта: " + data[4])
                f.write("\nПримеси: " + data[6])
                f.write("\nТолщина стенки: " + data[7])
                f.write("\nПоказатель водопоглощения: " + data[8])
                f.write("\nСледы брака: " + data[9])
                f.write("\nТекстура массы: " + data[10])
                f.write("\nМесто обнаружения: " + data[12]+"\n\n\n")
            f.close()
            os.startfile(r'artefact.txt')
        elif searchrol == '':
            self.lineEdit.setText("")
            self.lineEdit.setPlaceholderText("Введите ID артефакта")
            self.lineEdit.setStyleSheet("color: red;\n")
        else:
            self.lineEdit.setText("")
            self.lineEdit.setPlaceholderText("Артефакта с таким ID не существует")
        cur.close()
    def openWindow(self, Dialog):
        searchrol = str(self.lineEdit.text())
        cur = conn.cursor()
        cur.execute("SELECT * from artefacts WHERE id_artefact=?", (searchrol,))
        dat = cur.fetchall()
        if dat and searchrol != '':
            pdf = FPDF()
            for data in dat:
                pdf.add_page()
                pdf.add_font('sysfont', '', r"c:\WINDOWS\Fonts\arial.ttf", uni=True)
                pdf.set_font('sysfont', '', 18)
                pdf.cell(200, 40, txt="Отчет по артефакту, где ID "+ data[2], ln=1, align="C")
                pdf.set_font('sysfont', '', 14)
                pdf.line(40, 40, 180, 40)
                pdf.ln(90)
                pdf.cell(20, 15, txt="ID артефакта: " + data[2], ln=1, align="L")
                pdf.cell(20, 15, txt="Тип артефакта: " + data[5], ln=1, align="L")
                pdf.cell(20, 15, txt="Цвет артефакта: " + data[4], ln=1, align="L")
                pdf.cell(20, 15, txt="Примеси: " + data[6], ln=1, align="L")
                pdf.cell(20, 15, txt="Толщина стенки: " + data[7], ln=1, align="L")
                pdf.cell(20, 15, txt="Показатель водопоглощения: " + data[8], ln=1, align="L")
                pdf.cell(20, 15, txt="Следы брака: " + data[9], ln=1, align="L")
                pdf.cell(20, 15, txt="Текстура массы: " + data[10], ln=1, align="L")
            pdf.output("artefact.pdf", "F")
            os.startfile(r'artefact.pdf')
        elif searchrol == '':
            self.lineEdit.setText("")
            self.lineEdit.setPlaceholderText("Введите ID артефакта")
            self.lineEdit.setStyleSheet("color: red;\n")
        else:
            self.lineEdit.setText("")
            self.lineEdit.setPlaceholderText("Артефакта с таким ID не существует")
        cur.close()
    def closeWindow(self, Dialog):
        Dialog.hide()
    def search(self):
        self.window = QtWidgets.QMainWindow()
        self.ui = search.Ui_MainWindow()
        self.ui.setupUi(self.window)
        self.window.show()
    def setupUi(self, Dialog):
        Dialog.setObjectName("Dialog")
        Dialog.resize(600, 400)
        Dialog.setStyleSheet(" \n"
                                  "  font-size: 18px;\n")
        self.textEdit = QTextEdit()
        self.verticalLayout = QtWidgets.QVBoxLayout(Dialog)
        self.verticalLayout.setObjectName("verticalLayout")
        self.horizontalLayout = QtWidgets.QHBoxLayout()
        self.horizontalLayout.setObjectName("horizontalLayout")
        self.label = QtWidgets.QLabel(Dialog)
        font = QtGui.QFont()
        font.setPointSize(15)
        self.label.setFont(font)
        self.label.setObjectName("label")
        self.verticalLayout.addWidget(self.label)
        self.lineEdit = QtWidgets.QLineEdit(Dialog)
        self.lineEdit.setMinimumSize(QtCore.QSize(0, 60))
        font = QtGui.QFont()
        font.setPointSize(10)
        self.lineEdit.setFont(font)
        self.lineEdit.setObjectName("lineEdit")
        self.verticalLayout.addWidget(self.lineEdit)
        self.pushButton = QtWidgets.QPushButton(Dialog)
        self.pushButton.setMinimumSize(QtCore.QSize(0, 60))
        self.pushButton.setStyleSheet(" QPushButton#pushButton {\n"
                                        "     background-color: #EEDDFF;\n"
                                        "  border: 1px solid #7922CC;\n"
                                        "  border-radius: 25px;\n"
                                        "}\n"
                                        "QPushButton#pushButton:pressed {\n"
                                        "     background-color: rgb(224, 0, 0);     \n"
                                        " }\n"
                                        " QPushButton#pushButton:hover {\n"
                                        "     background-color:#FFE8DB;\n"
                                        " }\n"
                                        "")
        self.pushButton.setObjectName("pushButton")
        self.pushButton.clicked.connect(lambda: self.openWindow(Dialog))
        self.horizontalLayout.addWidget(self.pushButton)
        self.pushButton_3 = QtWidgets.QPushButton(Dialog)
        self.pushButton_3.setMinimumSize(QtCore.QSize(0, 60))
        self.pushButton_3.setStyleSheet(" QPushButton#pushButton {\n"
                                        "     background-color: #EEDDFF;\n"
                                        "  border: 1px solid #7922CC;\n"
                                        "  border-radius: 25px;\n"
                                        "}\n"
                                        "QPushButton#pushButton:pressed {\n"
                                        "     background-color: rgb(224, 0, 0);     \n"
                                        " }\n"
                                        " QPushButton#pushButton:hover {\n"
                                        "     background-color:#FFE8DB;\n"
                                        " }\n"
                                        "")
        self.pushButton_3.setObjectName("pushButton")
        self.pushButton_3.clicked.connect(lambda: self.openWindowtxt(Dialog))
        self.horizontalLayout.addWidget(self.pushButton_3)
        self.verticalLayout.addLayout(self.horizontalLayout)
        self.pushButton_2 = QtWidgets.QPushButton(Dialog)
        self.pushButton_2.setMinimumSize(QtCore.QSize(0, 60))
        self.pushButton_2.setStyleSheet(" QPushButton#pushButton_2 {\n"
                                        "     background-color: #EEDDFF;\n"
                                        "  border: 1px solid #7922CC;\n"
                                        "  border-radius: 25px;\n"
                                        "}\n"
                                        "QPushButton#pushButton_2:pressed {\n"
                                        "     background-color: rgb(224, 0, 0);     \n"
                                        " }\n"
                                        " QPushButton#pushButton_2:hover {\n"
                                        "     background-color:#FFE8DB;\n"
                                        " }\n"
                                        "")
        self.pushButton_2.setObjectName("pushButton_2")
        self.pushButton_2.clicked.connect(self.search)
        self.verticalLayout.addWidget(self.pushButton_2)
        self.pushButton_2.clicked.connect(lambda: self.closeWindow(Dialog))
        self.retranslateUi(Dialog)
        QtCore.QMetaObject.connectSlotsByName(Dialog)

    def retranslateUi(self, Dialog):
        Dialog.setWindowTitle(QtWidgets.QApplication.translate("Dialog", "Поиск по ID артефакта", None, -1))
        self.label.setText(QtWidgets.QApplication.translate("Dialog", "Введите ID "
"артефакта", None, -1))
        self.pushButton.setText(QtWidgets.QApplication.translate("Dialog", "Открыть в pdf", None, -1))
        self.pushButton_3.setText(QtWidgets.QApplication.translate("Dialog", "Открыть в txt", None, -1))
        self.pushButton_2.setText(QtWidgets.QApplication.translate("Dialog", "Отмена", None, -1))


if __name__ == "__main__":
    import sys
    app = QtWidgets.QApplication(sys.argv)
    Dialog = QtWidgets.QDialog()
    ui = Ui_Dialog()
    ui.setupUi(Dialog)
    Dialog.show()
    sys.exit(app.exec_())

